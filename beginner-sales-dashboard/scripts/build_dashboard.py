import random, datetime as dt
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter

random.seed(42); np.random.seed(42)
FONT = "Arial"
NAVY, TEAL, LIGHT, GREY = "1F3A5F", "1B998B", "E8F3F1", "F4F6F8"

# ---------------------------------------------------------------- data
PRODUCTS = [  # category, product, price, cost, popularity, max_qty
    ("Furniture", "Office Chair", 8500, 5600, .6, 2),
    ("Furniture", "Standing Desk", 19500, 13800, .35, 2),
    ("Furniture", "Bookshelf", 6200, 4100, .5, 3),
    ("Technology", "Wireless Mouse", 1200, 650, 1.2, 5),
    ("Technology", "Mechanical Keyboard", 4200, 2500, .9, 3),
    ("Technology", 'Monitor 24"', 11500, 8200, .6, 2),
    ("Office Supplies", "Notebook Set", 450, 190, 1.5, 8),
    ("Office Supplies", "Printer Paper Box", 1600, 1050, 1.3, 6),
    ("Office Supplies", "Desk Organizer", 850, 420, 1.0, 5),
    ("Accessories", "USB-C Hub", 1800, 980, 1.0, 4),
    ("Accessories", "Webcam", 3200, 1900, .7, 3),
    ("Accessories", "Headphones", 4500, 2600, .8, 3),
]
REGIONS = ["North", "South", "East", "West"]; RW = [.24, .31, .17, .28]
MONTH_W = [.8, .7, .9, .95, 1.0, .95, .9, 1.0, 1.1, 1.2, 1.5, 1.6]
pw = np.array([p[4] for p in PRODUCTS]); pw = pw / pw.sum()
mw = np.array(MONTH_W); mw = mw / mw.sum()

N = 1200
rows = []
for i in range(N):
    m = np.random.choice(12) if False else np.random.choice(range(1, 13), p=mw)
    d = dt.date(2025, m, random.randint(1, 28))
    cat, prod, price, cost, _, mq = PRODUCTS[np.random.choice(len(PRODUCTS), p=pw)]
    qty = random.randint(1, mq)
    disc = random.choice([0, 0, 0, 0, .05, .10, .15, .20])
    rows.append([f"ORD-{10001+i}", d, random.choices(REGIONS, RW)[0], cat, prod, qty, price, disc, cost])
cols = ["Order ID", "Order Date", "Region", "Category", "Product", "Quantity", "Unit Price", "Discount", "Unit Cost"]
base = pd.DataFrame(rows, columns=cols)

# ---- make the RAW version messy
raw = base.copy()
raw["Order Date"] = raw["Order Date"].astype(object)
idx = raw.sample(frac=.25, random_state=1).index
for i in idx:
    raw.at[i, "Order Date"] = raw.at[i, "Order Date"].strftime("%d-%b-%Y")   # text dates
var = {"North": ["north", " North ", "NORTH"], "South": ["south", "South ", "SOUTH"],
       "East": ["east", " East", "EAST"], "West": ["west", "West ", "WEST"]}
for i in raw.sample(n=140, random_state=2).index:
    raw.at[i, "Region"] = random.choice(var[raw.at[i, "Region"]])
for i in raw.sample(n=15, random_state=3).index:
    raw.at[i, "Unit Price"] = np.nan
for i in raw.sample(n=6, random_state=4).index:
    raw.at[i, "Quantity"] = np.nan
dups = raw.sample(n=30, random_state=5)
raw = pd.concat([raw, dups]).sample(frac=1, random_state=6).reset_index(drop=True)

# ---------------------------------------------------------------- cleaning (mirrors log)
log = []
c = raw.copy()
n0 = len(c)
log.append(("Raw records loaded", n0, "Starting point: 1,236 rows incl. duplicates and blanks".replace("1,236", f"{n0:,}")))
before = len(c); c = c.drop_duplicates(); log.append(("Exact duplicate rows removed", before - len(c), "Same Order ID and identical values"))
before = len(c); c = c.dropna(subset=["Quantity"]); log.append(("Rows dropped: missing Quantity", before - len(c), "Sales cannot be computed without a quantity"))
miss = int(c["Unit Price"].isna().sum())
price_map = c.groupby("Product")["Unit Price"].agg(lambda s: s.mode().iloc[0])
c["Unit Price"] = c["Unit Price"].fillna(c["Product"].map(price_map))
log.append(("Missing Unit Price filled", miss, "Filled with the product's standard (most common) price"))
before_reg = c["Region"].copy()
c["Region"] = c["Region"].str.strip().str.title()
log.append(("Region labels standardised", int((before_reg != c["Region"]).sum()), "Trimmed spaces and fixed upper/lower case"))
txt = c["Order Date"].apply(lambda v: isinstance(v, str))
log.append(("Text dates converted to real dates", int(txt.sum()), "Format DD-Mon-YYYY converted to Excel dates"))
c["Order Date"] = pd.to_datetime(c["Order Date"], format="mixed", dayfirst=True)
c["Quantity"] = c["Quantity"].astype(int)
c = c.sort_values(["Order Date", "Order ID"]).reset_index(drop=True)
log.append(("Final clean records", len(c), "Rows used in all analysis"))

c["Sales"] = c["Quantity"] * c["Unit Price"] * (1 - c["Discount"])
c["Profit"] = c["Sales"] - c["Quantity"] * c["Unit Cost"]
c["M"] = c["Order Date"].dt.month

# ---------------------------------------------------------------- workbook helpers
def f(bold=False, size=10, color="000000", italic=False):
    return Font(name=FONT, bold=bold, size=size, color=color, italic=italic)
def fill(hex_): return PatternFill("solid", start_color=hex_, end_color=hex_)
thin = Side(style="thin", color="D0D5DA"); box = Border(top=thin, bottom=thin, left=thin, right=thin)
INR = '"₹"#,##0'; PCT = "0.0%"

def header(ws, row, labels, col=1):
    for j, t in enumerate(labels):
        cell = ws.cell(row=row, column=col + j, value=t)
        cell.font = f(True, 10, "FFFFFF"); cell.fill = fill(NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = box

wb = Workbook()
dash = wb.active; dash.title = "Dashboard"
an = wb.create_sheet("Analysis"); ins = wb.create_sheet("Insights")
cd = wb.create_sheet("Clean Data"); rw = wb.create_sheet("Raw Data")

# ---------------------------------------------------------------- Raw Data
header(rw, 1, cols)
for r, row in enumerate(raw.itertuples(index=False), start=2):
    for j, v in enumerate(row, start=1):
        if isinstance(v, float) and np.isnan(v): v = None
        if isinstance(v, pd.Timestamp): v = v.to_pydatetime()
        if isinstance(v, dt.date) and not isinstance(v, dt.datetime): v = dt.datetime(v.year, v.month, v.day)
        cell = rw.cell(row=r, column=j, value=v); cell.font = f()
        if j == 2 and not isinstance(v, str): cell.number_format = "DD-MMM-YYYY"
        if j == 8: cell.number_format = "0%"
for j, w in enumerate([12, 14, 10, 16, 22, 10, 11, 10, 11], start=1):
    rw.column_dimensions[get_column_letter(j)].width = w
rw.freeze_panes = "A2"

# ---------------------------------------------------------------- Clean Data
ch = ["Order ID", "Order Date", "Month", "Region", "Category", "Product", "Quantity", "Unit Price", "Discount", "Unit Cost", "Sales", "Profit"]
header(cd, 1, ch)
for r, row in enumerate(c.itertuples(index=False), start=2):
    vals = {1: row[0], 2: row[1].to_pydatetime(), 4: row[2], 5: row[3], 6: row[4],
            7: int(row[5]), 8: float(row[6]), 9: float(row[7]), 10: float(row[8])}
    for j, v in vals.items():
        cd.cell(row=r, column=j, value=v).font = f()
    cd.cell(row=r, column=3, value=f"=MONTH(B{r})").font = f()
    cd.cell(row=r, column=11, value=f"=G{r}*H{r}*(1-I{r})").font = f()
    cd.cell(row=r, column=12, value=f"=K{r}-G{r}*J{r}").font = f()
    cd.cell(row=r, column=2).number_format = "DD-MMM-YYYY"
    cd.cell(row=r, column=8).number_format = INR; cd.cell(row=r, column=10).number_format = INR
    cd.cell(row=r, column=9).number_format = "0%"
    cd.cell(row=r, column=11).number_format = INR; cd.cell(row=r, column=12).number_format = INR
LAST = len(c) + 1
for j, w in enumerate([12, 14, 8, 10, 16, 22, 10, 12, 10, 12, 13, 13], start=1):
    cd.column_dimensions[get_column_letter(j)].width = w
cd.freeze_panes = "A2"; cd.auto_filter.ref = f"A1:L{LAST}"

def rng(col): return f"'Clean Data'!${col}$2:${col}${LAST}"

# ---------------------------------------------------------------- Analysis
an["A1"] = "Sales Analysis (all figures are formulas on the Clean Data sheet)"; an["A1"].font = f(True, 14, NAVY)
an["A2"] = "Currency: Indian Rupees (₹). Profit = Sales − Quantity × Unit Cost. Margin = Profit ÷ Sales."; an["A2"].font = f(False, 9, "666666", True)

# Monthly
an["A3"] = "1. Monthly trend"; an["A3"].font = f(True, 11, TEAL)
header(an, 4, ["Month", "Month #", "Orders", "Units", "Sales", "Profit", "Margin", "MoM Sales Growth"])
mn = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
for i in range(12):
    r = 5 + i
    an.cell(r, 1, mn[i]); an.cell(r, 2, i + 1)
    an.cell(r, 3, f"=COUNTIFS({rng('C')},B{r})")
    an.cell(r, 4, f"=SUMIFS({rng('G')},{rng('C')},B{r})")
    an.cell(r, 5, f"=SUMIFS({rng('K')},{rng('C')},B{r})")
    an.cell(r, 6, f"=SUMIFS({rng('L')},{rng('C')},B{r})")
    an.cell(r, 7, f"=IFERROR(F{r}/E{r},0)")
    an.cell(r, 8, "-" if i == 0 else f"=IFERROR(E{r}/E{r-1}-1,0)")
r = 17
an.cell(r, 1, "Total")
for col, L in [(3, "C"), (4, "D"), (5, "E"), (6, "F")]:
    an.cell(r, col, f"=SUM({L}5:{L}16)")
an.cell(r, 7, "=IFERROR(F17/E17,0)")

# Category
an["A20"] = "2. Category performance"; an["A20"].font = f(True, 11, TEAL)
header(an, 21, ["Category", "Orders", "Units", "Sales", "Profit", "Margin", "Share of Sales"])
cats = ["Technology", "Furniture", "Accessories", "Office Supplies"]
cat_sales = c.groupby("Category")["Sales"].sum().sort_values(ascending=False)
cats = list(cat_sales.index)
for i, k in enumerate(cats):
    r = 22 + i
    an.cell(r, 1, k)
    an.cell(r, 2, f"=COUNTIFS({rng('E')},A{r})")
    an.cell(r, 3, f"=SUMIFS({rng('G')},{rng('E')},A{r})")
    an.cell(r, 4, f"=SUMIFS({rng('K')},{rng('E')},A{r})")
    an.cell(r, 5, f"=SUMIFS({rng('L')},{rng('E')},A{r})")
    an.cell(r, 6, f"=IFERROR(E{r}/D{r},0)")
    an.cell(r, 7, f"=IFERROR(D{r}/$D$26,0)")
an.cell(26, 1, "Total")
for col, L in [(2, "B"), (3, "C"), (4, "D"), (5, "E")]:
    an.cell(26, col, f"=SUM({L}22:{L}25)")
an.cell(26, 6, "=IFERROR(E26/D26,0)"); an.cell(26, 7, "=SUM(G22:G25)")

# Product
an["A29"] = "3. Product performance (sorted by sales)"; an["A29"].font = f(True, 11, TEAL)
header(an, 30, ["Product", "Category", "Units", "Sales", "Profit", "Margin", "Sales Rank"])
prod_sales = c.groupby("Product")["Sales"].sum().sort_values(ascending=False)
pcat = {p[1]: p[0] for p in PRODUCTS}
for i, p in enumerate(prod_sales.index):
    r = 31 + i
    an.cell(r, 1, p); an.cell(r, 2, pcat[p])
    an.cell(r, 3, f"=SUMIFS({rng('G')},{rng('F')},A{r})")
    an.cell(r, 4, f"=SUMIFS({rng('K')},{rng('F')},A{r})")
    an.cell(r, 5, f"=SUMIFS({rng('L')},{rng('F')},A{r})")
    an.cell(r, 6, f"=IFERROR(E{r}/D{r},0)")
    an.cell(r, 7, f"=RANK(D{r},$D$31:$D$42)")

# Region
an["A45"] = "4. Regional performance"; an["A45"].font = f(True, 11, TEAL)
header(an, 46, ["Region", "Orders", "Sales", "Profit", "Margin", "Share of Sales"])
reg_sales = c.groupby("Region")["Sales"].sum().sort_values(ascending=False)
for i, k in enumerate(reg_sales.index):
    r = 47 + i
    an.cell(r, 1, k)
    an.cell(r, 2, f"=COUNTIFS({rng('D')},A{r})")
    an.cell(r, 3, f"=SUMIFS({rng('K')},{rng('D')},A{r})")
    an.cell(r, 4, f"=SUMIFS({rng('L')},{rng('D')},A{r})")
    an.cell(r, 5, f"=IFERROR(D{r}/C{r},0)")
    an.cell(r, 6, f"=IFERROR(C{r}/$C$51,0)")
an.cell(51, 1, "Total")
for col, L in [(2, "B"), (3, "C"), (4, "D")]:
    an.cell(51, col, f"=SUM({L}47:{L}50)")
an.cell(51, 5, "=IFERROR(D51/C51,0)"); an.cell(51, 6, "=SUM(F47:F50)")

# style analysis tables
def style_block(r0, r1, c0, c1, money_cols=(), pct_cols=(), int_cols=(), total_row=None):
    for r in range(r0, r1 + 1):
        for cc in range(c0, c1 + 1):
            cell = an.cell(r, cc); cell.font = f(r == total_row); cell.border = box
            if r == total_row: cell.fill = fill(LIGHT)
            if cc in money_cols: cell.number_format = INR
            elif cc in pct_cols: cell.number_format = PCT
            elif cc in int_cols: cell.number_format = "#,##0"
            if cc > c0 and r >= r0 and not isinstance(cell.value, str) or (isinstance(cell.value, str) and cell.value.startswith("=")):
                cell.alignment = Alignment(horizontal="right")
style_block(5, 17, 1, 8, money_cols=(5, 6), pct_cols=(7, 8), int_cols=(2, 3, 4), total_row=17)
style_block(22, 26, 1, 7, money_cols=(4, 5), pct_cols=(6, 7), int_cols=(2, 3), total_row=26)
style_block(31, 42, 1, 7, money_cols=(4, 5), pct_cols=(6,), int_cols=(3, 7))
style_block(47, 51, 1, 6, money_cols=(3, 4), pct_cols=(5, 6), int_cols=(2,), total_row=51)
for r in range(5, 17): an.cell(r, 2).alignment = Alignment(horizontal="center")
an.cell(5, 8).alignment = Alignment(horizontal="right")
for j, w in enumerate([20, 18, 12, 15, 15, 12, 16, 18], start=1):
    an.column_dimensions[get_column_letter(j)].width = w
an.row_dimensions[4].height = 30
an.freeze_panes = "A4"

# ---------------------------------------------------------------- computed findings (for text)
tot_sales = c["Sales"].sum(); tot_profit = c["Profit"].sum()
msales = c.groupby("M")["Sales"].sum()
peak_m, low_m = mn[msales.idxmax() - 1], mn[msales.idxmin() - 1]
nd_share = (msales[11] + msales[12]) / tot_sales
cat_p = c.groupby("Category")[["Sales", "Profit"]].sum(); cat_p["Margin"] = cat_p["Profit"] / cat_p["Sales"]
prod_p = c.groupby("Product")[["Sales", "Profit"]].sum(); prod_p["Margin"] = prod_p["Profit"] / prod_p["Sales"]
top_prod = prod_sales.index[0]; bot_prod = prod_sales.index[-1]
best_margin_cat = cat_p["Margin"].idxmax(); worst_margin_cat = cat_p["Margin"].idxmin()
best_margin_prod = prod_p["Margin"].idxmax(); worst_margin_prod = prod_p["Margin"].idxmin()
reg_share = reg_sales / tot_sales
top3_share = prod_sales.iloc[:3].sum() / tot_sales
disc_m = c.assign(D=c["Discount"] > 0).groupby("D")[["Sales", "Profit"]].sum()
disc_margin = disc_m.loc[True, "Profit"] / disc_m.loc[True, "Sales"]; nodisc_margin = disc_m.loc[False, "Profit"] / disc_m.loc[False, "Sales"]
def inr(x): return f"₹{x/1e5:,.1f} lakh" if x >= 1e5 else f"₹{x:,.0f}"

findings = [
    f"Sales peak in {peak_m} and are lowest in {low_m}; November and December together deliver {nd_share:.0%} of annual sales, so the business is strongly seasonal.",
    f"{top_prod} is the top product by sales ({inr(prod_sales.iloc[0])}); the top 3 products contribute {top3_share:.0%} of total sales. {bot_prod} is the weakest ({inr(prod_sales.iloc[-1])}).",
    f"{cat_sales.index[0]} is the largest category by sales ({cat_sales.iloc[0]/tot_sales:.0%} share), while {best_margin_cat} earns the best margin ({cat_p.loc[best_margin_cat,'Margin']:.1%}) and {worst_margin_cat} the lowest ({cat_p.loc[worst_margin_cat,'Margin']:.1%}).",
    f"{reg_sales.index[0]} is the strongest region ({reg_share.iloc[0]:.0%} of sales) and {reg_sales.index[-1]} the weakest ({reg_share.iloc[-1]:.0%}).",
    f"Discounted orders earn a {disc_margin:.1%} margin versus {nodisc_margin:.1%} for full-price orders, so discounting has a real cost to profitability.",
]
recs = [
    f"Build inventory and staffing ahead of Q4: stock the best-selling products before October to avoid missing the Nov–Dec peak.",
    f"Run targeted promotions in the slow months ({low_m} and neighbouring months) to smooth demand instead of discounting during peak season.",
    f"Protect margin on {worst_margin_prod} ({prod_p.loc[worst_margin_prod,'Margin']:.1%} margin) through price review or supplier negotiation; push {best_margin_prod} ({prod_p.loc[best_margin_prod,'Margin']:.1%} margin) in bundles.",
    f"Investigate why {reg_sales.index[-1]} lags: check local marketing, stock availability and delivery times before adding discounts.",
    "Set a discount cap for low-margin categories so promotions do not push orders below the target margin.",
]

# ---------------------------------------------------------------- Insights sheet
ins.column_dimensions["A"].width = 4; ins.column_dimensions["B"].width = 34; ins.column_dimensions["C"].width = 16; ins.column_dimensions["D"].width = 70
ins["B1"] = "Insights, Method and Data Notes"; ins["B1"].font = f(True, 14, NAVY)
r = 3
def section(title):
    global r
    ins.cell(r, 2, title).font = f(True, 11, TEAL); r += 1
def para(text, bullet=True):
    global r
    ins.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    cell = ins.cell(r, 2, ("•  " if bullet else "") + text); cell.font = f(); cell.alignment = Alignment(wrap_text=True, vertical="top")
    ins.row_dimensions[r].height = 15 * max(1, -(-len(text) // 105)) + 3; r += 1

section("About the dataset")
para("Retail office-products sales, January to December 2025, 4 regions, 4 categories, 12 products.")
para("IMPORTANT: this is a synthetic (simulated) dataset generated for this exercise, with deliberately messy raw data. Replace it with a real dataset (for example a public Superstore or Kaggle sales file) if your submission needs real-world data.")
r += 1
section("Data cleaning log")
header(ins, r, ["Step", "Rows affected", "Notes"], col=2); r += 1
for step, n, note in log:
    for j, v in enumerate([step, n, note], start=2):
        cell = ins.cell(r, j, v); cell.font = f(); cell.border = box
        cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if j == 3 else "left")
    r += 1
r += 1
section("Metrics chosen and why")
para("Sales = Quantity × Unit Price × (1 − Discount): the revenue actually collected per order.")
para("Profit and Margin: show which products earn money, not just which sell most.")
para("Orders and Average Order Value: separate demand volume from basket size.")
para("Month-over-month growth and share of sales: reveal seasonality and concentration.")
r += 1
section("Dashboard structure")
para("Top row: six headline KPIs. Below: monthly trend (time), category (group), product (detail), region (geography). Each chart answers one business question and reads from the Analysis sheet, which reads from Clean Data via formulas.")
r += 1
section("Key findings")
for t in findings: para(t)
r += 1
section("Recommendations")
for t in recs: para(t)

# ---------------------------------------------------------------- Dashboard
dash.sheet_view.showGridLines = False
for j in range(1, 15): dash.column_dimensions[get_column_letter(j)].width = 11
dash.column_dimensions["A"].width = 2; dash.column_dimensions["N"].width = 2
dash.merge_cells("B1:M1"); dash["B1"] = "Retail Sales Dashboard — 2025"
dash["B1"].font = f(True, 20, "FFFFFF"); dash["B1"].fill = fill(NAVY); dash["B1"].alignment = Alignment(vertical="center", indent=1)
dash.merge_cells("B2:M2"); dash["B2"] = "Beginner-level analysis: 12 months, 4 regions, 4 categories, 12 products (synthetic dataset). All numbers update from the Clean Data sheet."
dash["B2"].font = f(False, 9, "666666", True)
dash.row_dimensions[1].height = 34
kpis = [
    ("TOTAL SALES", "=Analysis!E17", INR),
    ("TOTAL PROFIT", "=Analysis!F17", INR),
    ("PROFIT MARGIN", "=Analysis!G17", PCT),
    ("ORDERS", "=Analysis!C17", "#,##0"),
    ("AVG ORDER VALUE", "=IFERROR(Analysis!E17/Analysis!C17,0)", INR),
    ("PEAK MONTH", "=INDEX(Analysis!A5:A16,MATCH(MAX(Analysis!E5:E16),Analysis!E5:E16,0))", "@"),
]
for i, (lab, fm, nf) in enumerate(kpis):
    c0 = 2 + i * 2
    dash.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c0 + 1)
    dash.merge_cells(start_row=5, start_column=c0, end_row=5, end_column=c0 + 1)
    a = dash.cell(4, c0, lab); a.font = f(True, 8, "555555"); a.alignment = Alignment(horizontal="center", vertical="bottom")
    v = dash.cell(5, c0, fm); v.font = f(True, 16, NAVY); v.number_format = nf; v.alignment = Alignment(horizontal="center", vertical="center")
    for rr in (4, 5):
        for cc in (c0, c0 + 1): dash.cell(rr, cc).fill = fill(LIGHT)
dash.row_dimensions[4].height = 20; dash.row_dimensions[5].height = 34

def style_chart(ch, title, w=12.6, h=7.6):
    ch.title = title; ch.title.overlay = False; ch.width = w; ch.height = h
    ch.x_axis.delete = False; ch.y_axis.delete = False
    ch.legend = None

# Monthly line
lc = LineChart(); style_chart(lc, "Monthly sales (₹)")
lc.add_data(Reference(an, min_col=5, min_row=4, max_row=16), titles_from_data=True)
lc.set_categories(Reference(an, min_col=1, min_row=5, max_row=16))
s = lc.series[0]; s.graphicalProperties.line.solidFill = TEAL; s.graphicalProperties.line.width = 28000; s.smooth = False
s.marker.symbol = "circle"; s.marker.size = 6; s.marker.graphicalProperties.solidFill = TEAL
lc.y_axis.numFmt = '#,##0'; lc.y_axis.majorGridlines.spPr = None
dash.add_chart(lc, "B7")

# Category bar
bc = BarChart(); bc.type = "col"; style_chart(bc, "Sales by category (₹)")
bc.add_data(Reference(an, min_col=4, min_row=21, max_row=25), titles_from_data=True)
bc.set_categories(Reference(an, min_col=1, min_row=22, max_row=25))
bc.series[0].graphicalProperties.solidFill = NAVY; bc.gapWidth = 60
bc.y_axis.numFmt = '#,##0'
dash.add_chart(bc, "H7")

# Product bar (horizontal)
pc = BarChart(); pc.type = "bar"; style_chart(pc, "Sales by product (₹)", h=9.4)
pc.add_data(Reference(an, min_col=4, min_row=30, max_row=42), titles_from_data=True)
pc.set_categories(Reference(an, min_col=1, min_row=31, max_row=42))
pc.series[0].graphicalProperties.solidFill = TEAL; pc.gapWidth = 40
pc.x_axis.scaling.orientation = "maxMin"; pc.y_axis.crosses = "max"; pc.y_axis.numFmt = "#,##0"
dash.add_chart(pc, "B24")

# Region bar
rc = BarChart(); rc.type = "col"; style_chart(rc, "Sales by region (₹)", h=9.4)
rc.add_data(Reference(an, min_col=3, min_row=46, max_row=50), titles_from_data=True)
rc.set_categories(Reference(an, min_col=1, min_row=47, max_row=50))
rc.series[0].graphicalProperties.solidFill = "E07A5F"; rc.gapWidth = 60
rc.y_axis.numFmt = '#,##0'
dash.add_chart(rc, "H24")

# findings on dashboard
fr = 44
dash.merge_cells(start_row=fr, start_column=2, end_row=fr, end_column=13)
dash.cell(fr, 2, "Key findings").font = f(True, 12, TEAL)
for i, t in enumerate(findings):
    rr = fr + 1 + i
    dash.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=13)
    cell = dash.cell(rr, 2, "•  " + t); cell.font = f(False, 10); cell.alignment = Alignment(wrap_text=True, vertical="top")
    dash.row_dimensions[rr].height = 30
dash.merge_cells(start_row=fr + 7, start_column=2, end_row=fr + 7, end_column=13)
dash.cell(fr + 7, 2, "See the Insights sheet for cleaning steps, metric choices and recommendations.").font = f(False, 9, "666666", True)

dash.sheet_properties.tabColor = NAVY; an.sheet_properties.tabColor = TEAL
wb.save("output/Beginner_Sales_Dashboard.xlsx")
print("saved", len(c), "clean rows;", "sales", round(tot_sales), "profit", round(tot_profit))
print(log)
for t in findings: print("-", t)

# export raw and cleaned data as CSV for the repo
raw.to_csv("data/raw_sales.csv", index=False)
c.drop(columns=["M"]).to_csv("data/clean_sales.csv", index=False)
