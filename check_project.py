"""Quick sanity tests for the beginner-sales-dashboard project.

Run from the project root (with the venv active):
    python check_project.py
It reads the files in data/ and output/ and prints PASS / FAIL for each check.
"""
import os
import sys
import pandas as pd
from openpyxl import load_workbook

RAW = "data/raw_sales.csv"
CLEAN = "data/clean_sales.csv"
XLSX = "output/Beginner_Sales_Dashboard.xlsx"

results = []


def check(name, ok, detail=""):
    results.append(ok)
    print(("PASS  " if ok else "FAIL  ") + name + (f"  ({detail})" if detail else ""))


# 1. files exist
for f in (RAW, CLEAN, XLSX):
    check(f"file exists: {f}", os.path.exists(f))
if not all(os.path.exists(f) for f in (RAW, CLEAN, XLSX)):
    sys.exit("Missing files - run 'python scripts/build_dashboard.py' from the project root first.")

raw = pd.read_csv(RAW)
cl = pd.read_csv(CLEAN)

# 2. raw data really is messy (the cleaning step has something to do)
check("raw: 1,230 rows", len(raw) == 1230, len(raw))
check("raw: 30 exact duplicate rows", raw.duplicated().sum() == 30, raw.duplicated().sum())
check("raw: 6 missing Quantity", raw["Quantity"].isna().sum() == 6, raw["Quantity"].isna().sum())
check("raw: 15 missing Unit Price", raw["Unit Price"].isna().sum() == 15, raw["Unit Price"].isna().sum())
check("raw: region labels are inconsistent", raw["Region"].nunique() > 4, raw["Region"].nunique())

# 3. clean data is actually clean
check("clean: 1,194 rows", len(cl) == 1194, len(cl))
check("clean: no missing values", cl.isna().sum().sum() == 0)
check("clean: no duplicate Order IDs", not cl["Order ID"].duplicated().any())
check("clean: exactly 4 tidy regions", sorted(cl["Region"].unique()) == ["East", "North", "South", "West"])
dates = pd.to_datetime(cl["Order Date"], errors="coerce")
check("clean: all dates valid and in 2025", dates.notna().all() and (dates.dt.year == 2025).all())
check("clean: Quantity >= 1 and Unit Price > 0", (cl["Quantity"] >= 1).all() and (cl["Unit Price"] > 0).all())

# 4. metrics recomputed independently of Excel
sales = cl["Quantity"] * cl["Unit Price"] * (1 - cl["Discount"])
profit = sales - cl["Quantity"] * cl["Unit Cost"]
check("Sales column matches formula", (abs(sales - cl["Sales"]) < 0.01).all())
check("Profit column matches formula", (abs(profit - cl["Profit"]) < 0.01).all())
check("total sales = 8,052,225", round(sales.sum()) == 8052225, round(sales.sum()))
check("total profit = 2,668,465", round(profit.sum()) == 2668465, round(profit.sum()))
check("profit margin about 33.1%", abs(profit.sum() / sales.sum() - 0.3314) < 0.001)
monthly = sales.groupby(dates.dt.month).sum()
check("peak month is December", monthly.idxmax() == 12, monthly.idxmax())

# 5. workbook structure
wb = load_workbook(XLSX)
check("workbook has 5 sheets", wb.sheetnames == ["Dashboard", "Analysis", "Insights", "Clean Data", "Raw Data"], wb.sheetnames)
check("Dashboard has 4 charts", len(wb["Dashboard"]._charts) == 4, len(wb["Dashboard"]._charts))
n_formulas = sum(1 for ws in wb for row in ws.iter_rows() for c in row
                 if isinstance(c.value, str) and c.value.startswith("="))
check("workbook contains formulas (not pasted values)", n_formulas > 3000, n_formulas)
check("Clean Data sheet has 1,194 data rows", wb["Clean Data"].max_row == 1195, wb["Clean Data"].max_row)

# 6. formula errors (only visible if Excel has calculated and saved the file)
wv = load_workbook(XLSX, data_only=True)
errors, uncalculated = [], 0
for ws in wb:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                v = wv[ws.title][c.coordinate].value
                if v is None:
                    uncalculated += 1
                elif isinstance(v, str) and v.startswith("#"):
                    errors.append((ws.title, c.coordinate, v))
check("no #REF!/#NAME?/#DIV/0! errors", not errors, errors[:3] if errors else "")
if uncalculated:
    print(f"NOTE  {uncalculated} formulas have no saved result yet. That is normal for a freshly built file: "
          "open it in Excel and press Ctrl+S once, then re-run this script for a full error check.")

print(f"\n{sum(results)}/{len(results)} checks passed")
sys.exit(0 if all(results) else 1)
